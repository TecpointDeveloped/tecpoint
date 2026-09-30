import json
import sys
from pathlib import Path

import openpyxl


source = Path(sys.argv[1])
target = Path(sys.argv[2])
workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
sheet = workbook[workbook.sheetnames[0]]
rows = sheet.iter_rows(values_only=True)
headers = [str(value or "").strip() for value in next(rows)]
records = []

for row_number, values in enumerate(rows, start=2):
    record = {headers[index]: value for index, value in enumerate(values)}
    record["sourceRow"] = row_number
    records.append(record)

target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"sheet": sheet.title, "rows": len(records), "columns": headers}, ensure_ascii=False))
